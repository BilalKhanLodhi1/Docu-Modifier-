"""Template-driven DOCX formatting engine.

A template is a plain configuration dict (stored on DocumentTemplate.config),
so formatting is never hardcoded into processing functions. `apply_template`
always reads the source file and writes a NEW file.
"""
import copy
import logging
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

from .docx_utils import classify_paragraph, heading_level, iter_all_paragraphs

log = logging.getLogger("docucraft.format")

PAGE_SIZES = {"A4": (Mm(210), Mm(297)), "Letter": (Inches(8.5), Inches(11)), "Legal": (Inches(8.5), Inches(14))}

DEFAULT_CONFIG = {
    "page_size": "A4",
    "orientation": "portrait",
    "margins_in": {"top": 1.0, "bottom": 1.0, "left": 1.0, "right": 1.0},
    "body_font": "Times New Roman",
    "body_size": 12,
    "line_spacing": 1.15,
    "space_after_pt": 8,
    "justify": True,
    "first_line_indent_in": 0.0,
    "heading_font": "Times New Roman",
    "heading_sizes": {"1": 16, "2": 14, "3": 12},
    "title_size": 22,
    "heading_color": "000000",
    "text_color": "1A1A1A",
    "code_font": "Consolas",
    "code_size": 10,
    "code_shading": "F2F2F2",
    "page_break_before_h1": False,
    "cover": {"enabled": True, "style": "academic"},
    "toc": False,
    "header_text": "",
    "footer_text": "",
    "page_numbers": True,
    "page_number_style": "Page {PAGE}",
    "table_borders": True,
}


def merged_config(config):
    """Deep-merge a stored template config over DEFAULT_CONFIG."""
    out = copy.deepcopy(DEFAULT_CONFIG)
    for k, v in (config or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k].update(v)
        else:
            out[k] = v
    return out


def _hex(c, fallback="000000"):
    c = (c or fallback).lstrip("#")
    try:
        return RGBColor.from_string(c.upper()[:6])
    except ValueError:
        return RGBColor.from_string(fallback)


def _set_style_font(style, name, size=None, bold=None, color=None, italic=None):
    font = style.font
    font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    if color:
        font.color.rgb = _hex(color)


def _get_style(document, name, base_type=1):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, base_type)


def _clear_run_overrides(paragraph, keep_emphasis=True):
    """Remove direct font name/size/color on runs so the template styles apply."""
    for run in paragraph.runs:
        rpr = run._element.rPr
        if rpr is None:
            continue
        for tag in ("w:rFonts", "w:sz", "w:szCs", "w:color", "w:highlight"):
            for el in rpr.findall(qn(tag)):
                rpr.remove(el)
        if not keep_emphasis:
            for tag in ("w:b", "w:i"):
                for el in rpr.findall(qn(tag)):
                    rpr.remove(el)


def _shade(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    for old in ppr.findall(qn("w:shd")):
        ppr.remove(old)
    ppr.append(shd)


def _add_field(paragraph, instr, placeholder="1"):
    """Append a complex field (PAGE, NUMPAGES, TOC) to a paragraph."""
    def el(tag, **attrs):
        e = OxmlElement(tag)
        for k, v in attrs.items():
            e.set(qn(k), v)
        return e
    r1 = paragraph.add_run()._r
    r1.append(el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
    r2 = paragraph.add_run()._r
    it = el("w:instrText", **{"xml:space": "preserve"})
    it.text = f" {instr} "
    r2.append(it)
    r3 = paragraph.add_run()._r
    r3.append(el("w:fldChar", **{"w:fldCharType": "separate"}))
    paragraph.add_run(placeholder)
    r5 = paragraph.add_run()._r
    r5.append(el("w:fldChar", **{"w:fldCharType": "end"}))


def _setup_pages(document, cfg):
    w, h = PAGE_SIZES.get(cfg["page_size"], PAGE_SIZES["A4"])
    landscape = cfg.get("orientation") == "landscape"
    m = cfg["margins_in"]
    for s in document.sections:
        s.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = (h, w) if landscape else (w, h)
        s.top_margin, s.bottom_margin = Inches(float(m["top"])), Inches(float(m["bottom"]))
        s.left_margin, s.right_margin = Inches(float(m["left"])), Inches(float(m["right"]))


def _setup_styles(document, cfg):
    normal = document.styles["Normal"]
    _set_style_font(normal, cfg["body_font"], cfg["body_size"], color=cfg["text_color"])
    pf = normal.paragraph_format
    pf.line_spacing = float(cfg["line_spacing"])
    pf.space_after = Pt(float(cfg["space_after_pt"]))
    for lvl, size in cfg["heading_sizes"].items():
        st = _get_style(document, f"Heading {lvl}")
        _set_style_font(st, cfg["heading_font"], size, bold=True, color=cfg["heading_color"], italic=False)
        st.paragraph_format.space_before = Pt(14 if lvl == "1" else 10)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
    title = _get_style(document, "Title")
    _set_style_font(title, cfg["heading_font"], cfg["title_size"], bold=True, color=cfg["heading_color"])
    tppr = title.element.get_or_add_pPr()
    for el in tppr.findall(qn("w:pBdr")):
        tppr.remove(el)
    title.paragraph_format.space_after = Pt(12)
    for lvl in cfg["heading_sizes"]:
        hppr = _get_style(document, f"Heading {lvl}").element.get_or_add_pPr()
        for el in hppr.findall(qn("w:pBdr")):
            hppr.remove(el)
    cap = _get_style(document, "Caption")
    _set_style_font(cap, cfg["body_font"], max(9, cfg["body_size"] - 2), italic=True, color=cfg["text_color"], bold=False)
    code = _get_style(document, "DocuCraft Code")
    _set_style_font(code, cfg["code_font"], cfg["code_size"], color="000000")
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0


def _format_body(document, cfg, scope_ids=None):
    """Apply paragraph-level formatting. scope_ids limits it to chosen top-level sections."""
    in_scope = scope_ids is None
    prev_kind = None
    first_h1 = True
    for idx, (p, table) in enumerate(iter_all_paragraphs(document)):
        kind = classify_paragraph(p)
        if table is None and kind in ("heading", "title") and (heading_level(p) or 0) <= 1 and scope_ids is not None:
            in_scope = f"p{idx}" in scope_ids
        if not in_scope:
            continue
        pf = p.paragraph_format
        if kind == "code":
            p.style = document.styles["DocuCraft Code"]
            _clear_run_overrides(p)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.first_line_indent = None
            if cfg.get("code_shading"):
                _shade(p, cfg["code_shading"])
            pf.space_before = Pt(6) if prev_kind != "code" else Pt(0)
        elif kind in ("heading", "title"):
            _clear_run_overrides(p)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if kind == "heading" else WD_ALIGN_PARAGRAPH.CENTER
            if (kind == "heading" and heading_level(p) == 1 and cfg.get("page_break_before_h1")
                    and table is None and not first_h1):
                pf.page_break_before = True
            if kind == "heading" and heading_level(p) == 1:
                first_h1 = False
        elif kind == "caption":
            p.style = document.styles["Caption"]
            _clear_run_overrides(p)
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif kind in ("paragraph", "list"):
            _clear_run_overrides(p)
            if table is None:
                if "<w:drawing" in p._p.xml:
                    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif kind == "paragraph" and cfg.get("justify"):
                    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    if cfg.get("first_line_indent_in"):
                        pf.first_line_indent = Inches(float(cfg["first_line_indent_in"]))
                pf.line_spacing = float(cfg["line_spacing"])
                pf.space_after = Pt(float(cfg["space_after_pt"]))
        elif kind == "empty" and "<w:drawing" in p._p.xml and table is None:
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        prev_kind = kind


def _new_paragraph_before(anchor_el, document):
    p = document.add_paragraph()
    anchor_el.addprevious(p._p)
    return p


def _build_cover(document, cfg, meta):
    """Insert a cover page as its own vertically-centred section at the start of the body."""
    body = document.element.body
    first = body[0]
    style = (cfg.get("cover") or {}).get("style", "academic")
    font = cfg["heading_font"]
    created = []

    cover_style = _ensure_paragraph_style(document, COVER_STYLE)

    def line(text, size=12, bold=False, italic=False, space_after=4, color=None):
        p = _new_paragraph_before(first, document)
        p.style = cover_style
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.first_line_indent = None
        if text:
            r = p.add_run(text)
            r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
            if color:
                r.font.color.rgb = _hex(color)
        created.append(p)
        return p

    logo = meta.get("logo_path")
    if logo and Path(logo).exists():
        p = line("", space_after=14)
        try:
            p.add_run().add_picture(str(logo), width=Inches(1.5))
        except Exception:
            log.warning("Logo could not be inserted")
    if style == "corporate":
        if meta.get("organization"):
            line(meta["organization"], 13, bold=True, space_after=2, color="333333")
        if meta.get("department"):
            line(meta["department"], 11, space_after=40, color="555555")
        line(meta.get("title") or "Untitled document", 26, bold=True, space_after=8)
        if meta.get("subtitle"):
            line(meta["subtitle"], 14, italic=True, space_after=40, color="444444")
        rows = [("Prepared for", meta.get("submitted_to")), ("Prepared by", meta.get("author")),
                ("Reference", meta.get("report_number")), ("Date", meta.get("date"))]
        for label, val in rows:
            if val:
                p = line("", 11, space_after=2)
                r = p.add_run(f"{label}: ")
                r.font.bold, r.font.size, r.font.name = True, Pt(11), font
                r2 = p.add_run(val)
                r2.font.size, r2.font.name = Pt(11), font
    else:
        if meta.get("organization"):
            line(meta["organization"], 18, bold=True, space_after=4)
        if meta.get("department"):
            line(meta["department"], 14, space_after=4)
        if meta.get("program"):
            line(meta["program"], 13, space_after=4)
        if meta.get("course"):
            line(meta["course"], 13, italic=True, space_after=4)
        line("", space_after=26)
        if meta.get("report_number"):
            line(meta["report_number"], 14, bold=True, space_after=8)
        line(meta.get("title") or "Untitled document", 24, bold=True, space_after=8)
        if meta.get("subtitle"):
            line(meta["subtitle"], 14, italic=True, space_after=8)
        line("", space_after=30)
        if meta.get("author") or meta.get("author_id"):
            line("Submitted by", 12, bold=True, space_after=2)
            if meta.get("author"):
                line(meta["author"], 13, space_after=2)
            if meta.get("author_id"):
                line(meta["author_id"], 12, space_after=12)
        if meta.get("submitted_to"):
            line("Submitted to", 12, bold=True, space_after=2)
            line(meta["submitted_to"], 13, space_after=12)
        if meta.get("date"):
            line(meta["date"], 12, space_after=0)

    # End the cover with a section break: copy the document's sectPr into the
    # last cover paragraph, centre it vertically, and drop headers/footers.
    body_sectpr = body.find(qn("w:sectPr"))
    cover_sectpr = copy.deepcopy(body_sectpr) if body_sectpr is not None else OxmlElement("w:sectPr")
    for tag in ("w:headerReference", "w:footerReference", "w:titlePg", "w:pgNumType", "w:vAlign"):
        for el in cover_sectpr.findall(qn(tag)):
            cover_sectpr.remove(el)
    last = created[-1] if created else line("")
    last._p.get_or_add_pPr().append(cover_sectpr)
    # Vertical centring that renders identically in Word and LibreOffice:
    # estimate the cover height and push the first line down.
    est = 0.0
    for p in created:
        sizes = [r.font.size.pt for r in p.runs if r.font.size] or [12]
        est += max(sizes) * 1.25 + (p.paragraph_format.space_after.pt if p.paragraph_format.space_after else 0)
        if "<w:drawing" in p._p.xml:
            est += 110
    sec = document.sections[0]
    avail = (sec.page_height - sec.top_margin - sec.bottom_margin) / 12700 if sec.page_height else 698
    m = cfg["margins_in"]
    page_h = PAGE_SIZES.get(cfg["page_size"], PAGE_SIZES["A4"])[1] / 12700
    avail = page_h - (float(m["top"]) + float(m["bottom"])) * 72
    top = max(0, (avail - est) / 2 - 24)
    if created:
        created[0].paragraph_format.space_before = Pt(top)
    # Restart numbering at 1 for the body section.
    if body_sectpr is not None:
        for el in body_sectpr.findall(qn("w:pgNumType")):
            body_sectpr.remove(el)
        pg = OxmlElement("w:pgNumType")
        pg.set(qn("w:start"), "1")
        body_sectpr.append(pg)
    return len(created)


COVER_STYLE = "DocuCraft Cover"
TOC_HEADING_STYLE = "TOC Heading"
TOC_LEVEL_STYLES = ("TOC 1", "TOC 2", "TOC 3")
GENERATED_STYLES = {COVER_STYLE, TOC_HEADING_STYLE, *TOC_LEVEL_STYLES}
BOOKMARK_PREFIX = "_TocDC"


def _ensure_paragraph_style(document, name, base="Normal"):
    try:
        return document.styles[name]
    except KeyError:
        st = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        try:
            st.base_style = document.styles[base]
        except KeyError:
            pass
        return st


def _remove_generated_front_matter(document):
    """Remove a cover page, contents page and heading bookmarks that DocuCraft added earlier.

    Without this, formatting an already-formatted version stacks a second cover and contents page.
    """
    body = document.element.body
    removed = 0
    for p in list(body.iterchildren(qn("w:p"))):
        ps = p.find(qn("w:pPr"))
        sid = ps.find(qn("w:pStyle")).get(qn("w:val")) if ps is not None and ps.find(qn("w:pStyle")) is not None else ""
        try:
            name = document.styles.get_by_id(sid, 1).name if sid else ""
        except Exception:
            name = ""
        if name in GENERATED_STYLES:
            body.remove(p)
            removed += 1
    ids = set()
    for bm in list(body.iter(qn("w:bookmarkStart"))):
        if (bm.get(qn("w:name")) or "").startswith(BOOKMARK_PREFIX):
            ids.add(bm.get(qn("w:id")))
            bm.getparent().remove(bm)
    for be in list(body.iter(qn("w:bookmarkEnd"))):
        if be.get(qn("w:id")) in ids:
            be.getparent().remove(be)
    return removed


def _text_width_in(cfg):
    w, h = PAGE_SIZES.get(cfg["page_size"], PAGE_SIZES["A4"])
    width = h if cfg.get("orientation") == "landscape" else w
    m = cfg["margins_in"]
    return width / 914400 - float(m["left"]) - float(m["right"])


def _setup_toc_styles(document, cfg):
    width = Inches(max(2.0, _text_width_in(cfg) - 0.02))
    head = _ensure_paragraph_style(document, TOC_HEADING_STYLE)
    head.font.name, head.font.size, head.font.bold = cfg["heading_font"], Pt(int((cfg.get("heading_sizes") or {}).get("1", 16))), True
    head.font.color.rgb = _hex(cfg.get("heading_color", "000000"))
    head.paragraph_format.space_before, head.paragraph_format.space_after = Pt(0), Pt(14)
    head.paragraph_format.first_line_indent = None
    for lvl, name in enumerate(TOC_LEVEL_STYLES, start=1):
        st = _ensure_paragraph_style(document, name)
        st.font.name, st.font.size = cfg["body_font"], Pt(cfg["body_size"])
        st.font.bold = False
        pf = st.paragraph_format
        pf.left_indent = Inches(0.28 * (lvl - 1))
        pf.first_line_indent = None
        pf.space_before = Pt(6 if lvl == 1 else 0)
        pf.space_after = Pt(2)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)


def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def _insert_toc(document, cfg, after_paragraphs):
    """Insert a contents page as its own section, with linked entries and a live Word TOC field.

    Each listed heading gets a bookmark; each entry links to it. Page numbers are left blank here
    and filled in by fill_toc_page_numbers() from a real layout pass, so they are correct in PDF
    exports and LibreOffice as well as in Word (which can still refresh the field after edits).
    """
    body = document.element.body
    anchor = body[after_paragraphs] if after_paragraphs < len(body) else None
    headings = [p for p, t in iter_all_paragraphs(document)
                if t is None and classify_paragraph(p) == "heading" and (heading_level(p) or 9) <= 3]
    _setup_toc_styles(document, cfg)

    def add(p):
        if anchor is not None:
            anchor.addprevious(p._p)
        return p

    for i, hp in enumerate(headings):
        name = f"{BOOKMARK_PREFIX}{i:04d}"
        start, end = _el("w:bookmarkStart", **{"w:id": str(90000 + i), "w:name": name}), _el("w:bookmarkEnd", **{"w:id": str(90000 + i)})
        ppr = hp._p.find(qn("w:pPr"))
        (ppr.addnext(start) if ppr is not None else hp._p.insert(0, start))
        hp._p.append(end)

    head = add(document.add_paragraph("Contents"))
    head.style = document.styles[TOC_HEADING_STYLE]
    entries = []
    for i, hp in enumerate(headings):
        lvl = min(3, max(1, heading_level(hp) or 1))
        p = add(document.add_paragraph())
        p.style = document.styles[TOC_LEVEL_STYLES[lvl - 1]]
        if i == 0:
            p.add_run()._r.append(_el("w:fldChar", **{"w:fldCharType": "begin", "w:dirty": "true"}))
            it = _el("w:instrText", **{"xml:space": "preserve"})
            it.text = ' TOC \\o "1-3" \\h \\z \\u '
            p.add_run()._r.append(it)
            p.add_run()._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
        link = _el("w:hyperlink", **{"w:anchor": f"{BOOKMARK_PREFIX}{i:04d}", "w:history": "1"})
        r = p.add_run(" ".join(hp.text.split()))
        link.append(r._r)
        p._p.append(link)
        p.add_run("\t")
        p.add_run("")  # page number, filled in by fill_toc_page_numbers()
        entries.append(p)
    if entries:
        entries[-1].add_run()._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    else:
        p = add(document.add_paragraph("No headings were found. Apply Heading 1 to 3 styles in Word to list sections here."))
        p.style = document.styles[TOC_LEVEL_STYLES[0]]
        entries.append(p)

    # End the contents page with its own section: roman numerals, and the body restarts at 1.
    body_sectpr = body.find(qn("w:sectPr"))
    toc_sectpr = copy.deepcopy(body_sectpr) if body_sectpr is not None else OxmlElement("w:sectPr")
    for tag in ("w:headerReference", "w:footerReference", "w:titlePg", "w:pgNumType", "w:vAlign", "w:type"):
        for el in toc_sectpr.findall(qn(tag)):
            toc_sectpr.remove(el)
    toc_sectpr.append(_el("w:pgNumType", **{"w:fmt": "lowerRoman", "w:start": "1"}))
    entries[-1]._p.get_or_add_pPr().append(toc_sectpr)
    _restart_body_numbering(document)

    settings_el = document.settings.element
    if settings_el.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings_el.append(uf)
    return len(headings)


def _restart_body_numbering(document):
    body_sectpr = document.element.body.find(qn("w:sectPr"))
    if body_sectpr is None:
        return
    for el in body_sectpr.findall(qn("w:pgNumType")):
        body_sectpr.remove(el)
    body_sectpr.append(_el("w:pgNumType", **{"w:start": "1"}))


def fill_toc_page_numbers(path):
    """Write real page numbers into the contents page by laying the document out once.

    A throwaway copy gets tiny invisible markers at each listed heading and at the start of the body;
    LibreOffice turns it into a PDF and the markers' pages give exact body page numbers.
    Returns True when numbers were filled in, False if LibreOffice is unavailable.
    """
    import tempfile
    from .conversion_service import docx_to_pdf, find_soffice
    if not find_soffice():
        return False
    import pymupdf

    document = docx.Document(str(path))
    entries = [p for p in document.paragraphs if p.style is not None and p.style.name in TOC_LEVEL_STYLES]
    if not entries or not any(p._p.find(qn("w:hyperlink")) is not None for p in entries):
        return False

    with tempfile.TemporaryDirectory(prefix="docucraft_toc_") as tmp:
        probe = docx.Document(str(path))
        body = probe.element.body

        def mark(par_el, label):
            r = _el("w:r")
            rpr = _el("w:rPr")
            rpr.append(_el("w:sz", **{"w:val": "2"}))
            r.append(rpr)
            t = _el("w:t")
            t.text = label
            r.append(t)
            ppr = par_el.find(qn("w:pPr"))
            (ppr.addnext(r) if ppr is not None else par_el.insert(0, r))

        for bm in body.iter(qn("w:bookmarkStart")):
            name = bm.get(qn("w:name")) or ""
            if name.startswith(BOOKMARK_PREFIX):
                par = bm.getparent()
                while par is not None and par.tag != qn("w:p"):
                    par = par.getparent()
                if par is not None:
                    mark(par, f"QQ{name[len(BOOKMARK_PREFIX):]}QQ")
        # The body starts at the first paragraph after the contents section break.
        toc_end = None
        for p in probe.paragraphs:
            if p.style is not None and p.style.name in TOC_LEVEL_STYLES and p._p.find(qn("w:pPr")) is not None \
                    and p._p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
                toc_end = p._p
        if toc_end is None:
            return False
        nxt = toc_end.getnext()
        while nxt is not None and nxt.tag != qn("w:p"):
            inner = next(nxt.iter(qn("w:p")), None) if nxt.tag == qn("w:tbl") else None
            if inner is not None:
                nxt = inner
                break
            nxt = nxt.getnext()
        if nxt is None:
            return False
        mark(nxt, "QQBODYQQ")
        src, pdf = Path(tmp) / "probe.docx", Path(tmp) / "probe.pdf"
        probe.save(str(src))
        docx_to_pdf(src, pdf, timeout=180)
        pages = {}
        with pymupdf.open(str(pdf)) as d:
            for i, page in enumerate(d):
                text = page.get_text().replace(" ", "").replace("\n", "")
                for token in __import__("re").findall(r"QQ(BODY|\d{4})QQ", text):
                    pages.setdefault(token, i)
    if "BODY" not in pages:
        return False
    for p in entries:
        link = p._p.find(qn("w:hyperlink"))
        if link is None:
            continue
        key = (link.get(qn("w:anchor")) or "")[len(BOOKMARK_PREFIX):]
        if key in pages:
            # The number goes in the last plain run; on the final entry the run after it closes the field.
            slots = [r for r in p.runs if not r._r.findall(qn("w:fldChar")) and not r._r.findall(qn("w:instrText"))]
            if slots:
                slots[-1].text = str(pages[key] - pages["BODY"] + 1)
    document.save(str(path))
    return True


def _header_footer(document, cfg):
    body_section = document.sections[-1]
    for s in document.sections:
        s.header.is_linked_to_previous = False
        s.footer.is_linked_to_previous = False
    hdr, ftr = body_section.header, body_section.footer
    has_front = len(document.sections) > 1
    if cfg.get("header_text"):
        p = hdr.paragraphs[0] if hdr.paragraphs else hdr.add_paragraph()
        p.text = cfg["header_text"]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in p.runs:
            r.font.size, r.font.name = Pt(9), cfg["body_font"]
            r.font.color.rgb = _hex("555555")
    if cfg.get("page_numbers") or cfg.get("footer_text"):
        p = ftr.paragraphs[0] if ftr.paragraphs else ftr.add_paragraph()
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if cfg.get("footer_text"):
            r = p.add_run(cfg["footer_text"] + ("    " if cfg.get("page_numbers") else ""))
            r.font.size = Pt(9)
        if cfg.get("page_numbers"):
            pattern = cfg.get("page_number_style") or "{PAGE}"
            if has_front:
                # Front matter is numbered separately, so "of N" must count body pages only.
                pattern = pattern.replace("{NUMPAGES}", "{SECTIONPAGES}")
            for i, chunk in enumerate(pattern.replace("{SECTIONPAGES}", "\x00SECTIONPAGES\x00").replace("{NUMPAGES}", "\x00NUMPAGES\x00").replace("{PAGE}", "\x00PAGE\x00").split("\x00")):
                if chunk in ("PAGE", "NUMPAGES", "SECTIONPAGES"):
                    _add_field(p, chunk)
                elif chunk:
                    p.add_run(chunk)
            for r in p.runs:
                r.font.size, r.font.name = Pt(9), cfg["body_font"]
    # Cover section (if any) gets empty header/footer
    if len(document.sections) > 1:
        first = document.sections[0]
        for part in (first.header, first.footer):
            for para in part.paragraphs:
                for r in list(para.runs):
                    r._element.getparent().remove(r._element)


def _style_tables(document, cfg=None):
    """Clean grid tables: shaded, bold, repeating header row; padded cells; space after the table."""
    try:
        grid = document.styles["Table Grid"]
    except KeyError:
        grid = None
    fill = ((cfg or {}).get("table_header_fill") or "E8ECF2").lstrip("#")
    for table in document.tables:
        sname = (table.style.name if table.style is not None else "") or ""
        if grid is not None and sname in ("", "Normal Table"):
            table.style = grid
        tblpr = table._tbl.tblPr
        if tblpr.find(qn("w:tblCellMar")) is None:
            mar = _el("w:tblCellMar")
            for side, w in (("top", "40"), ("left", "100"), ("bottom", "40"), ("right", "100")):
                mar.append(_el(f"w:{side}", **{"w:w": w, "w:type": "dxa"}))
            tblpr.append(mar)
        if not table.rows:
            continue
        header = table.rows[0]
        trpr = header._tr.get_or_add_trPr()
        if trpr.find(qn("w:tblHeader")) is None:
            trpr.append(_el("w:tblHeader"))
        for cell in header.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            if tcpr.find(qn("w:shd")) is None:
                tcpr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.bold = True
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.first_line_indent = None
        nxt = table._tbl.getnext()
        if nxt is not None and nxt.tag == qn("w:p"):
            from docx.text.paragraph import Paragraph
            para = Paragraph(nxt, table._parent)
            if not para.paragraph_format.space_before or para.paragraph_format.space_before < Pt(8):
                para.paragraph_format.space_before = Pt(10)


def heading_outline(path):
    """Top-level headings with stable ids, used for 'apply to selected sections'."""
    d = docx.Document(str(path))
    out = []
    for idx, (p, t) in enumerate(iter_all_paragraphs(d)):
        if t is None and classify_paragraph(p) in ("heading", "title") and (heading_level(p) or 0) <= 1:
            out.append({"id": f"p{idx}", "text": p.text.strip()[:120]})
    return out


def apply_template(src_path, dst_path, config=None, metadata=None, scope_ids=None,
                   include_cover=None, include_toc=None):
    cfg = merged_config(config)
    meta = metadata or {}
    document = docx.Document(str(src_path))
    _remove_generated_front_matter(document)
    _setup_styles(document, cfg)
    _format_body(document, cfg, scope_ids=scope_ids)
    if cfg.get("table_borders"):
        _style_tables(document, cfg)
    cover_on = cfg["cover"].get("enabled") if include_cover is None else include_cover
    toc_on = cfg.get("toc") if include_toc is None else include_toc
    cover_len = 0
    if cover_on:
        cover_len = _build_cover(document, cfg, meta)
    if toc_on:
        _insert_toc(document, cfg, cover_len)
    _setup_pages(document, cfg)
    _header_footer(document, cfg)
    if toc_on and cfg.get("page_numbers"):
        toc_section = document.sections[-2]
        p = toc_section.footer.paragraphs[0] if toc_section.footer.paragraphs else toc_section.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(p, "PAGE", "i")
        for r in p.runs:
            r.font.size, r.font.name = Pt(9), cfg["body_font"]
    if cover_on and not toc_on:
        _restart_body_numbering(document)
    if meta.get("title"):
        document.core_properties.title = meta["title"][:250]
    if meta.get("author"):
        document.core_properties.author = meta["author"][:250]
    document.save(str(dst_path))
    numbered = False
    if toc_on:
        try:
            numbered = fill_toc_page_numbers(dst_path)
        except Exception:
            log.warning("Could not compute contents page numbers; Word will fill them in on open", exc_info=True)
    return {"cover": bool(cover_on), "toc": bool(toc_on), "toc_page_numbers": numbered, "scoped": scope_ids is not None}
