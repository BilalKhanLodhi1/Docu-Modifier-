"""Extract, enhance and reinsert DOCX images.

Core operations are local (Pillow + OpenCV). Optional background removal uses
`rembg` when installed. Optional AI image editing goes through a pluggable
provider and is never applied without explicit user approval.
"""
import io
import logging

import docx
from docx.oxml.ns import qn
from PIL import Image, ImageEnhance, ImageFilter, ImageOps

from .docx_utils import iter_all_paragraphs
from .exceptions import DependencyMissingError, DocuCraftError

log = logging.getLogger("docucraft.images")

EMU_PER_INCH = 914400
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def _blips(paragraph):
    return paragraph._p.findall(f".//{{{A_NS}}}blip")


def list_images(path):
    """Return image descriptors in document order.

    Each item: index, rel_id, content_type, width/height px, display width
    (inches), paragraph index, nearby caption and surrounding text.
    """
    document = docx.Document(str(path))
    paras = list(iter_all_paragraphs(document))
    items = []
    for pidx, (p, table) in enumerate(paras):
        for blip in _blips(p):
            rid = blip.get(f"{{{R_NS}}}embed")
            if not rid or rid not in p.part.related_parts:
                continue
            part = p.part.related_parts[rid]
            blob = part.blob
            try:
                with Image.open(io.BytesIO(blob)) as im:
                    w, h, fmt = im.width, im.height, im.format
            except Exception:
                w = h = None
                fmt = part.content_type.split("/")[-1].upper()
            extent = blip.getparent()
            # walk up to wp:inline / wp:anchor to read display size
            node, cx = blip, None
            while node is not None:
                if node.tag in (f"{{{WP_NS}}}inline", f"{{{WP_NS}}}anchor"):
                    ext = node.find(f"{{{WP_NS}}}extent")
                    if ext is not None:
                        cx = int(ext.get("cx"))
                    break
                node = node.getparent()
            caption = ""
            for j in (pidx + 1, pidx - 1):
                if 0 <= j < len(paras):
                    t = paras[j][0].text.strip()
                    if t.lower().startswith(("figure", "fig.", "image", "screenshot", "chart")):
                        caption = t
                        break
            context = ""
            for j in range(pidx - 1, -1, -1):
                if paras[j][0].text.strip():
                    context = paras[j][0].text.strip()[:120]
                    break
            items.append({
                "index": len(items),
                "rel_id": rid,
                "paragraph_index": pidx,
                "content_type": part.content_type,
                "format": fmt,
                "width": w, "height": h,
                "size_bytes": len(blob),
                "display_width_in": round(cx / EMU_PER_INCH, 2) if cx else None,
                "dpi_estimate": round(w / (cx / EMU_PER_INCH)) if (w and cx) else None,
                "caption": caption,
                "context": context,
                "in_table": table is not None,
            })
    return items


def get_image_bytes(path, index):
    document = docx.Document(str(path))
    for i, (blip, part_owner) in enumerate(_iter_blips(document)):
        if i == index:
            rid = blip.get(f"{{{R_NS}}}embed")
            return part_owner.related_parts[rid].blob
    raise DocuCraftError("That image could not be found in the document.")


def _iter_blips(document):
    for p, _ in iter_all_paragraphs(document):
        for blip in _blips(p):
            rid = blip.get(f"{{{R_NS}}}embed")
            if rid and rid in p.part.related_parts:
                yield blip, p.part


# ----------------------------------------------------------------------------
# Enhancement operations
# ----------------------------------------------------------------------------
def process_image(data, ops):
    """Apply a dict of operations to image bytes. Returns (bytes, format, warnings).

    ops keys (all optional): crop {x,y,w,h in %}, rotate (90/180/270), flip_h, flip_v,
    resize_width (px, aspect preserved), brightness, contrast, sharpness, saturation
    (1.0 = unchanged), denoise (0-20), auto_contrast, grayscale, remove_background,
    format (PNG/JPEG/WEBP), quality (JPEG/WEBP 40-95), optimize.
    """
    warnings = []
    img = Image.open(io.BytesIO(data))
    img = ImageOps.exif_transpose(img)
    src_format = (img.format or "PNG").upper()
    if img.mode == "P":
        img = img.convert("RGBA")

    crop = ops.get("crop")
    if crop and (float(crop.get("x") or 0) > 0 or float(crop.get("y") or 0) > 0
                 or float(crop.get("w") or 100) < 100 or float(crop.get("h") or 100) < 100):
        x = max(0.0, min(99.0, float(crop.get("x", 0)))) / 100
        y = max(0.0, min(99.0, float(crop.get("y", 0)))) / 100
        w = max(1.0, min(100.0, float(crop.get("w", 100)))) / 100
        h = max(1.0, min(100.0, float(crop.get("h", 100)))) / 100
        box = (int(img.width * x), int(img.height * y),
               int(min(img.width, img.width * (x + w))), int(min(img.height, img.height * (y + h))))
        if box[2] - box[0] > 4 and box[3] - box[1] > 4:
            img = img.crop(box)
    rot = int(ops.get("rotate") or 0)
    if rot in (90, 180, 270):
        img = img.rotate(-rot, expand=True)
    if ops.get("flip_h"):
        img = ImageOps.mirror(img)
    if ops.get("flip_v"):
        img = ImageOps.flip(img)

    rw = int(ops.get("resize_width") or 0)
    if rw and rw != img.width:
        rw = max(16, min(rw, 8000))
        if rw > img.width * 1.5:
            warnings.append("Upscaling more than 1.5x can make the image look soft.")
        rh = max(1, round(img.height * rw / img.width))
        img = img.resize((rw, rh), Image.LANCZOS)

    if ops.get("remove_background"):
        try:
            from rembg import remove  # optional dependency
        except ImportError:
            raise DependencyMissingError("Background removal needs the optional 'rembg' package: pip install rembg")
        buf = io.BytesIO()
        img.save(buf, "PNG")
        img = Image.open(io.BytesIO(remove(buf.getvalue())))

    denoise = int(ops.get("denoise") or 0)
    if denoise > 0:
        img = _denoise(img, min(denoise, 20))

    has_alpha = img.mode in ("RGBA", "LA")
    alpha = img.getchannel("A") if has_alpha else None
    rgb = img.convert("RGB")
    if ops.get("auto_contrast"):
        rgb = ImageOps.autocontrast(rgb, cutoff=1)
    for key, enh in (("brightness", ImageEnhance.Brightness), ("contrast", ImageEnhance.Contrast),
                     ("saturation", ImageEnhance.Color)):
        v = float(ops.get(key) or 1.0)
        if abs(v - 1.0) > 0.01:
            rgb = enh(rgb).enhance(max(0.0, min(v, 3.0)))
    sharp = float(ops.get("sharpness") or 1.0)
    if sharp > 1.01:
        rgb = rgb.filter(ImageFilter.UnsharpMask(radius=2, percent=int(min(sharp, 3.0) * 70), threshold=2))
    elif sharp < 0.99:
        rgb = ImageEnhance.Sharpness(rgb).enhance(sharp)
    if ops.get("grayscale"):
        rgb = ImageOps.grayscale(rgb).convert("RGB")
    img = rgb
    if alpha is not None and alpha.size == rgb.size:
        img = rgb.convert("RGBA")
        img.putalpha(alpha)

    fmt = (ops.get("format") or src_format).upper()
    if fmt == "JPG":
        fmt = "JPEG"
    if fmt not in ("PNG", "JPEG", "WEBP", "GIF", "BMP", "TIFF"):
        fmt = "PNG"
    out = io.BytesIO()
    save_kwargs = {}
    if fmt in ("JPEG", "WEBP"):
        save_kwargs["quality"] = int(ops.get("quality") or 88)
        if img.mode in ("RGBA", "LA", "P"):
            if fmt == "JPEG":
                warnings.append("JPEG has no transparency; transparent areas were filled with white.")
                bg = Image.new("RGB", img.size, "white")
                bg.paste(img, mask=img.getchannel("A"))
                img = bg
    if ops.get("optimize"):
        save_kwargs["optimize"] = True
    img.save(out, fmt, **save_kwargs)
    return out.getvalue(), fmt, warnings


def _denoise(img, strength):
    try:
        import cv2
        import numpy as np
    except ImportError:
        return img.filter(ImageFilter.MedianFilter(3))
    mode = img.mode
    arr = np.array(img.convert("RGB"))[:, :, ::-1]
    arr = cv2.fastNlMeansDenoisingColored(arr, None, strength, strength, 7, 21)
    out = Image.fromarray(arr[:, :, ::-1])
    if mode in ("RGBA", "LA"):
        out = out.convert("RGBA")
        out.putalpha(img.getchannel("A"))
    return out


# ----------------------------------------------------------------------------
# Reinsertion
# ----------------------------------------------------------------------------
_CT = {"PNG": "image/png", "JPEG": "image/jpeg", "GIF": "image/gif", "BMP": "image/bmp",
       "TIFF": "image/tiff", "WEBP": "image/webp"}


def replace_image(src_path, dst_path, index, new_bytes, width_in=None, alignment=None, caption=None):
    """Replace image `index` in a copy of the document.

    The image keeps its position and display width (unless width_in is given);
    height is recalculated from the new pixel aspect ratio so it never distorts.
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.packuri import PackURI
    from docx.parts.image import ImagePart

    document = docx.Document(str(src_path))
    target = None
    for i, (pair, _table) in enumerate(_iter_para_blips(document)):
        if i == index:
            target = pair
            break
    if target is None:
        raise DocuCraftError("That image could not be found in the document.")
    paragraph, blip = target
    with Image.open(io.BytesIO(new_bytes)) as im:
        new_w, new_h, fmt = im.width, im.height, (im.format or "PNG").upper()
    if fmt == "WEBP":
        # Word does not reliably display WebP; store as PNG.
        with Image.open(io.BytesIO(new_bytes)) as im:
            buf = io.BytesIO()
            im.save(buf, "PNG")
            new_bytes, fmt = buf.getvalue(), "PNG"

    # Always add a NEW image part so images shared elsewhere are not changed.
    package = paragraph.part.package
    ext = {"JPEG": "jpg"}.get(fmt, fmt.lower())
    existing = {str(p.partname) for p in package.iter_parts()}
    n = 1
    while f"/word/media/docucraft_{n}.{ext}" in existing:
        n += 1
    part = ImagePart(PackURI(f"/word/media/docucraft_{n}.{ext}"), _CT.get(fmt, "image/png"), new_bytes, None)
    new_rid = paragraph.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
    blip.set(f"{{{R_NS}}}embed", new_rid)

    # Resize display box, preserving aspect ratio.
    node = blip
    while node is not None and node.tag not in (f"{{{WP_NS}}}inline", f"{{{WP_NS}}}anchor"):
        node = node.getparent()
    if node is not None:
        ext_el = node.find(f"{{{WP_NS}}}extent")
        cx = int(ext_el.get("cx"))
        if width_in:
            cx = int(float(width_in) * EMU_PER_INCH)
        section = document.sections[0]
        max_cx = int(section.page_width - section.left_margin - section.right_margin)
        cx = min(cx, max_cx)
        cy = int(cx * new_h / new_w)
        ext_el.set("cx", str(cx))
        ext_el.set("cy", str(cy))
        for x in node.iter(f"{{{A_NS}}}ext"):
            if x.get("cx") is not None:
                x.set("cx", str(cx))
                x.set("cy", str(cy))
    if alignment in ("left", "center", "right"):
        paragraph.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                               "right": WD_ALIGN_PARAGRAPH.RIGHT}[alignment]
    if caption is not None:
        _set_caption(document, paragraph, caption)
    document.save(str(dst_path))
    return {"width_px": new_w, "height_px": new_h, "format": fmt}


def _iter_para_blips(document):
    for p, t in iter_all_paragraphs(document):
        for blip in _blips(p):
            rid = blip.get(f"{{{R_NS}}}embed")
            if rid and rid in p.part.related_parts:
                yield (p, blip), t


def _set_caption(document, image_paragraph, caption):
    from docx.text.paragraph import Paragraph
    from .docx_utils import set_paragraph_text
    nxt = image_paragraph._p.getnext()
    if nxt is not None and nxt.tag == qn("w:p"):
        np_ = Paragraph(nxt, image_paragraph._parent)
        if np_.text.strip().lower().startswith(("figure", "fig.", "image", "chart", "screenshot")):
            if caption.strip():
                set_paragraph_text(np_, caption.strip())
            else:
                nxt.getparent().remove(nxt)
            return
    if caption.strip():
        new_p = document.add_paragraph(caption.strip())
        try:
            new_p.style = document.styles["Caption"]
        except KeyError:
            pass
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph._p.addnext(new_p._p)


# ----------------------------------------------------------------------------
# Optional AI image provider
# ----------------------------------------------------------------------------
class AIImageProvider:
    name = "none"

    def available(self):
        return False, "No AI image provider is configured (AI_IMAGE_PROVIDER=none)."

    def edit(self, image_bytes, instruction):
        raise DependencyMissingError("AI image editing is not configured. Local enhancements are still available.")


def get_image_provider():
    from django.conf import settings
    name = getattr(settings, "AI_IMAGE_PROVIDER", "none")
    # Extension point: register real providers here (e.g. a Stable Diffusion
    # img2img endpoint). Kept optional so core tools never depend on it.
    return AIImageProvider()
