"""Output quality: contents page, re-formatting, tables and formatting kept through rewrites."""
import unittest

import docx
from docx.oxml.ns import qn
from django.test import SimpleTestCase

from services.conversion_service import find_soffice
from services.docx_utils import set_paragraph_text
from services.docx_writer import apply_template, merged_config

from .base import DocuCraftTestCase


def long_report(path, sections=4):
    d = docx.Document()
    d.add_paragraph("Annual Review", style="Title")
    text = "The team recorded each request and compared monthly figures with the previous year. " * 4
    for i in range(1, sections + 1):
        d.add_paragraph(f"Section {i}", style="Heading 1")
        d.add_paragraph(f"Part {i}.1", style="Heading 2")
        for _ in range(7):
            d.add_paragraph(text)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Value", "A", "1"
    d.add_paragraph("After the table.")
    d.save(str(path))


class TemplateQualityTests(DocuCraftTestCase):
    def format(self, src, dst, cover=True, toc=True):
        cfg = merged_config({"page_numbers": True, "page_number_style": "Page {PAGE} of {NUMPAGES}"})
        return apply_template(src, dst, cfg, {"title": "Annual Review", "author": "Tester"},
                              include_cover=cover, include_toc=toc)

    def test_formatting_twice_does_not_stack_cover_or_contents(self):
        src, once, twice = self.fx / "long.docx", self.fx / "once.docx", self.fx / "twice.docx"
        long_report(src)
        self.format(src, once)
        self.format(once, twice)
        d = docx.Document(str(twice))
        styles = [p.style.name for p in d.paragraphs]
        self.assertEqual(styles.count("TOC Heading"), 1)
        self.assertEqual(len(d.sections), 3, "cover, contents and body")
        self.assertEqual(sum(s.startswith("TOC ") and s != "TOC Heading" for s in styles), 8)
        names = [b.get(qn("w:name")) for b in d.element.body.iter(qn("w:bookmarkStart"))]
        self.assertEqual(len(names), len(set(names)), "heading bookmarks must not be duplicated")

    def test_contents_entries_link_to_headings_and_field_is_well_formed(self):
        src, out = self.fx / "long.docx", self.fx / "out.docx"
        long_report(src)
        self.format(src, out)
        d = docx.Document(str(out))
        kinds = [f.get(qn("w:fldCharType")) for f in d.element.body.iter(qn("w:fldChar"))]
        self.assertEqual(kinds, ["begin", "separate", "end"])
        anchors = {h.get(qn("w:anchor")) for h in d.element.body.iter(qn("w:hyperlink"))}
        marks = {b.get(qn("w:name")) for b in d.element.body.iter(qn("w:bookmarkStart"))}
        self.assertTrue(anchors and anchors <= marks)
        self.assertNotIn("Contents", [p.text for p in d.paragraphs if p.style.name.startswith("Heading")])

    @unittest.skipUnless(find_soffice(), "LibreOffice is not installed")
    def test_contents_page_numbers_match_the_pages(self):
        import pymupdf
        from services.conversion_service import docx_to_pdf
        src, out, pdf = self.fx / "long.docx", self.fx / "out.docx", self.fx / "out.pdf"
        long_report(src, sections=5)
        result = self.format(src, out, cover=False)
        self.assertTrue(result["toc_page_numbers"])
        d = docx.Document(str(out))
        entries = [p.text.split("\t") for p in d.paragraphs if p.style.name in ("TOC 1", "TOC 2")]
        docx_to_pdf(out, pdf)
        with pymupdf.open(str(pdf)) as doc:
            pages = [[l.strip() for l in pg.get_text().split("\n")] for pg in doc]
        body_start = 1  # contents page first, body starts on the next page
        numbers = set()
        for name, number in entries:
            actual = next(i for i in range(body_start, len(pages)) if name in pages[i]) - body_start + 1
            self.assertEqual(int(number), actual, name)
            numbers.add(actual)
        self.assertGreater(len(numbers), 2, "the test document should span several pages")

    def test_tables_get_a_repeating_shaded_header_and_space_after(self):
        src, out = self.fx / "long.docx", self.fx / "out.docx"
        long_report(src, sections=1)
        self.format(src, out, cover=False, toc=False)
        d = docx.Document(str(out))
        header = d.tables[0].rows[0]
        self.assertIsNotNone(header._tr.trPr.find(qn("w:tblHeader")))
        self.assertIsNotNone(header.cells[0]._tc.tcPr.find(qn("w:shd")))
        after = next(p for p in d.paragraphs if p.text == "After the table.")
        self.assertGreaterEqual(after.paragraph_format.space_before.pt, 8)


class KeepFormattingTests(SimpleTestCase):
    def paragraph(self):
        d = docx.Document()
        p = d.add_paragraph()
        p.add_run("the ")
        p.add_run("BankAccount").bold = True
        p.add_run(" class store a ")
        p.add_run("initial").italic = True
        p.add_run(" balance and use ")
        p.add_run("deposit()").font.name = "Consolas"
        p.add_run(", see note")
        p.add_run("2").font.superscript = True
        return p

    def test_inline_formatting_survives_rewording(self):
        p = self.paragraph()
        set_paragraph_text(p, "The BankAccount class stores an initial balance and uses deposit(), see note 2")
        self.assertEqual(p.text, "The BankAccount class stores an initial balance and uses deposit(), see note 2")
        styled = {r.text: (r.bold, r.italic, r.font.name, r.font.superscript) for r in p.runs}
        self.assertTrue(styled["BankAccount"][0])
        self.assertTrue(styled["initial"][1])
        self.assertEqual(styled["deposit()"][2], "Consolas")
        self.assertTrue(styled["2"][3])

    def test_formatting_dropped_cleanly_when_word_is_gone(self):
        p = self.paragraph()
        set_paragraph_text(p, "The BankAccount class keeps a starting balance.")
        self.assertEqual(p.text, "The BankAccount class keeps a starting balance.")
        self.assertFalse(any(r.italic for r in p.runs))
